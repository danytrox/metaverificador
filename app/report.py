"""Exportación de resultados a CSV, JSON, HTML y Excel (todo local, sin CDN ni red)."""
from __future__ import annotations

import csv
import html
import json
from datetime import datetime
from typing import Iterable

from .analyzer import Summary
from .template import TemplateSpec, _as_text, is_author_column, resolve_cell


def _summaries_to_list(items: Iterable[Summary]) -> list[Summary]:
    return list(items)


def export_csv(
    summaries: Iterable[Summary],
    path: str,
    template: TemplateSpec | None = None,
    deleted: Iterable[Summary] | None = None,
) -> None:
    summaries = _summaries_to_list(summaries)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        if template is None:
            w.writerow([
                "Archivo", "Ruta", "Tipo", "Autor", "Autor_valores", "Titulo",
                "Fecha_creacion", "Software", "Advertencias",
            ])
            for s in summaries:
                w.writerow([
                    s.filename,
                    s.path,
                    s.filetype,
                    "SÍ" if s.author_found else "NO",
                    " | ".join(s.author_values),
                    s.title,
                    s.creation_date,
                    " | ".join(s.software),
                    " | ".join(s.warnings),
                ])
        else:
            w.writerow([c.header for c in template.columns])
            for s in summaries:
                w.writerow([resolve_cell(c, s) for c in template.columns])
        if deleted:
            deleted = list(deleted)
            w.writerow([])
            w.writerow(["Archivos eliminados (excluidos del análisis)", ""])
            w.writerow(["Archivo", "Estado"])
            for s in deleted:
                w.writerow([s.filename, "Eliminado por el usuario"])


def export_json(
    summaries: Iterable[Summary],
    path: str,
    deleted: Iterable[Summary] | None = None,
) -> None:
    summaries = _summaries_to_list(summaries)
    payload = []
    for s in summaries:
        rec = s.to_dict()
        rec["tags"] = s.tags
        payload.append(rec)
    doc = {
        "generated": datetime.now().isoformat(timespec="seconds"),
        "backend": "local",
        "files": payload,
    }
    if deleted:
        doc["deleted_files"] = [
            {"filename": s.filename, "path": s.path, "state": "eliminado"}
            for s in list(deleted)
        ]
    with open(path, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=2)


_DEFAULT_HTML_COLUMNS = ["Archivo", "Tipo", "Autor", "Autor (valores)", "Título", "Fecha", "Software", "Advertencias", "Detalle"]

_HTML_TEMPLATE = """<!doctype html>
<html lang="es">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>MetaVerificador - Reporte</title>
<style>
 body{{font-family:Segoe UI,Arial,sans-serif;margin:24px;color:#1a1a1a;}}
 h1{{font-size:20px;}}
 .meta{{color:#666;font-size:13px;margin-bottom:16px;}}
 table{{border-collapse:collapse;width:100%;font-size:13px;}}
 th,td{{border:1px solid #ddd;padding:6px 8px;text-align:left;vertical-align:top;}}
 th{{background:#f2f2f2;}}
 tr.sin-autor td{{background:#fdecea;}}
 .pill{{display:inline-block;padding:1px 8px;border-radius:10px;font-size:11px;font-weight:bold;}}
 .si{{background:#d4edda;color:#155724;}}
 .no{{background:#f8d7da;color:#721c24;}}
 details{{margin-top:4px;}}
 pre{{background:#f7f7f7;padding:8px;overflow:auto;font-size:11px;}}
</style>
</head>
<body>
<h1>MetaVerificador - Reporte de metadatos</h1>
<div class="meta">Generado {generated} · {n} archivo(s) · análisis 100% local</div>
<table>
<thead><tr>{headers}</tr></thead>
<tbody>
{rows}
</tbody>
</table>
{deleted_section}
</body>
</html>
"""


def export_html(
    summaries: Iterable[Summary],
    path: str,
    template: TemplateSpec | None = None,
    deleted: Iterable[Summary] | None = None,
) -> None:
    summaries = _summaries_to_list(summaries)
    rows = []
    if template is None:
        header_cells = "".join(f"<th>{html.escape(c)}</th>" for c in _DEFAULT_HTML_COLUMNS)
        for s in summaries:
            cls = "" if s.author_found else "sin-autor"
            pill = '<span class="pill si">SÍ</span>' if s.author_found else '<span class="pill no">NO</span>'
            tags_html = "<pre>" + html.escape(
                "\n".join(f"{k}: {v}" for k, v in sorted(s.tags.items()))
            ) + "</pre>"
            rows.append(
                f'<tr class="{cls}">'
                f'<td>{html.escape(s.filename)}</td>'
                f'<td>{html.escape(s.filetype)}</td>'
                f'<td>{pill}</td>'
                f'<td>{html.escape("; ".join(s.author_values))}</td>'
                f'<td>{html.escape(s.title)}</td>'
                f'<td>{html.escape(s.creation_date)}</td>'
                f'<td>{html.escape("; ".join(s.software))}</td>'
                f'<td>{html.escape("; ".join(s.warnings))}</td>'
                f'<td><details><summary>ver</summary>{tags_html}</details></td>'
                f'</tr>'
            )
    else:
        columns = template.columns
        header_cells = "".join(f"<th>{html.escape(c.header)}</th>" for c in columns)
        for s in summaries:
            cls = "" if s.author_found else "sin-autor"
            cells = []
            for c in columns:
                val = resolve_cell(c, s)
                if c.kind == "author_bool":
                    cells.append(
                        f'<td><span class="pill si">SÍ</span></td>'
                        if s.author_found else
                        f'<td><span class="pill no">NO</span></td>'
                    )
                elif c.kind == "tags":
                    cells.append(f'<td><details><summary>ver</summary><pre>{html.escape(val)}</pre></details></td>')
                else:
                    cells.append(f"<td>{html.escape(val)}</td>")
            rows.append(f'<tr class="{cls}">' + "".join(cells) + "</tr>")
    deleted_html = ""
    if deleted:
        deleted = list(deleted)
        items = "".join(f"<li>{html.escape(s.filename)}</li>" for s in deleted)
        deleted_html = (
            '<h2>Archivos eliminados</h2>'
            f'<div class="meta">Se quitaron {len(deleted)} archivo(s) de la lista '
            "antes de exportar:</div>"
            f"<ul>{items}</ul>"
        )
    doc = _HTML_TEMPLATE.format(
        generated=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        n=len(summaries),
        headers=header_cells,
        rows="\n".join(rows),
        deleted_section=deleted_html,
    )
    with open(path, "w", encoding="utf-8") as f:
        f.write(doc)


def _xlsx_styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    return {
        "header_fill": PatternFill("solid", fgColor="1F4E79"),
        "header_font": Font(bold=True, color="FFFFFF", size=11),
        "title_font": Font(bold=True, color="FFFFFF", size=14),
        "sub_font": Font(italic=True, color="595959", size=10),
        "legend_font": Font(italic=True, color="595959", size=9),
        "no_fill": PatternFill("solid", fgColor="FDE7E9"),
        "yes_font": Font(bold=True, color="1E7B34"),
        "no_font": Font(bold=True, color="C00000"),
        "border": Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        ),
        "center": Alignment(horizontal="center", vertical="center"),
        "top_wrap": Alignment(vertical="top", wrap_text=True),
    }


def _add_metadata_sheet(wb, summaries: list[Summary]) -> None:
    """Hoja 'Metadatos': volcado completo campo por campo."""
    from openpyxl.utils import get_column_letter

    st = _xlsx_styles()
    ws = wb.create_sheet("Metadatos")
    cols = ["Archivo", "Grupo", "Campo", "Valor"]
    for i, name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.font = st["header_font"]
        cell.fill = st["header_fill"]
        cell.alignment = st["center"]
        cell.border = st["border"]

    r = 2
    for s in summaries:
        for k in sorted(s.tags):
            grupo, _, campo = k.partition(":")
            vals = [s.filename, grupo, campo, _as_text(s.tags[k])]
            for i, val in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=i, value=val)
                cell.border = st["border"]
                cell.alignment = st["top_wrap"]
            r += 1

    for i, w in enumerate([38, 16, 30, 70], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    if r > 2:
        ws.auto_filter.ref = f"A1:D{r - 1}"


def _add_deleted_sheet(wb, deleted: list[Summary]) -> None:
    """Hoja 'Eliminados': archivos que el usuario quitó antes de exportar."""
    if not deleted:
        return
    from openpyxl.utils import get_column_letter

    st = _xlsx_styles()
    ws = wb.create_sheet("Eliminados")
    cols = ["Archivo", "Ruta", "Estado"]
    for i, name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.font = st["header_font"]
        cell.fill = st["header_fill"]
        cell.alignment = st["center"]
        cell.border = st["border"]

    for r, s in enumerate(deleted, start=2):
        for i, val in enumerate([s.filename, s.path, "Eliminado por el usuario"], start=1):
            cell = ws.cell(row=r, column=i, value=val)
            cell.border = st["border"]
            cell.alignment = st["top_wrap"]

    for i, w in enumerate([38, 70, 24], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def _add_template_sheet(ws, summaries: list[Summary], template: TemplateSpec) -> None:
    """Hoja 'Datos' generada a partir de una plantilla de columnas."""
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    st = _xlsx_styles()
    cols = template.columns
    last_col = get_column_letter(max(len(cols), 1))

    ws.merge_cells(f"A1:{last_col}1")
    c = ws["A1"]
    c.value = f"MetaVerificador — {template.name}"
    c.font = st["title_font"]
    c.fill = st["header_fill"]
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells(f"A2:{last_col}2")
    c = ws["A2"]
    c.value = (
        f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}   ·   "
        f"{len(summaries)} archivo(s)   ·   análisis 100% local"
    )
    c.font = st["sub_font"]
    ws.row_dimensions[2].height = 16

    header_row = 3
    for i, col in enumerate(cols, start=1):
        cell = ws.cell(row=header_row, column=i, value=col.header)
        cell.font = st["header_font"]
        cell.fill = st["header_fill"]
        cell.alignment = st["center"]
        cell.border = st["border"]

    for r, s in enumerate(summaries, start=header_row + 1):
        for i, col in enumerate(cols, start=1):
            val = resolve_cell(col, s)
            cell = ws.cell(row=r, column=i, value=val)
            cell.border = st["border"]
            cell.alignment = st["top_wrap"]
            if not s.author_found:
                cell.fill = st["no_fill"]
            if is_author_column(col):
                cell.font = st["yes_font"] if s.author_found else st["no_font"]
                cell.alignment = st["center"]

    for i in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 30
    ws.freeze_panes = f"B{header_row + 1}"
    last_row = header_row + len(summaries)
    if summaries:
        ws.auto_filter.ref = f"A{header_row}:{last_col}{last_row}"


def export_xlsx(
    summaries: Iterable[Summary],
    path: str,
    template: TemplateSpec | None = None,
    deleted: Iterable[Summary] | None = None,
) -> None:
    """Exporta a Excel (.xlsx).

    Sin plantilla: dos hojas — "Resumen" (una fila por archivo con autor SÍ/NO,
    filas rojas para los sin autor) y "Metadatos" (volcado completo).
    Con plantilla: hoja "Datos" con las columnas pedidas por la plantilla, más
    la hoja "Metadatos" con el volcado completo.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    summaries = _summaries_to_list(summaries)
    st = _xlsx_styles()

    wb = Workbook()

    if template is None:
        # ---------- Hoja 1: Resumen ----------
        ws = wb.active
        ws.title = "Resumen"
        cols = ["Archivo", "Ruta", "Tipo", "Autor", "Autor (detalle)",
                "Título", "Fecha de creación", "Software", "Advertencias"]
        last_col = get_column_letter(len(cols))

        ws.merge_cells(f"A1:{last_col}1")
        c = ws["A1"]
        c.value = "MetaVerificador — Reporte de metadatos"
        c.font = st["title_font"]
        c.fill = st["header_fill"]
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 26

        ws.merge_cells(f"A2:{last_col}2")
        c = ws["A2"]
        c.value = (
            f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}   ·   "
            f"{len(summaries)} archivo(s)   ·   análisis 100% local"
        )
        c.font = st["sub_font"]
        ws.row_dimensions[2].height = 16

        ws.merge_cells(f"A3:{last_col}3")
        c = ws["A3"]
        c.value = (
            "Leyenda:  SÍ = registra autor en metadatos  ·  NO = sin autor (fila en rojo)  ·  "
            "\"Autor (detalle)\" indica el campo exacto y el valor encontrado."
        )
        c.font = st["legend_font"]

        header_row = 4
        for i, name in enumerate(cols, start=1):
            cell = ws.cell(row=header_row, column=i, value=name)
            cell.font = st["header_font"]
            cell.fill = st["header_fill"]
            cell.alignment = st["center"]
            cell.border = st["border"]

        for r, s in enumerate(summaries, start=header_row + 1):
            vals = [
                s.filename,
                s.path,
                s.filetype,
                "SÍ" if s.author_found else "NO",
                "; ".join(s.author_values),
                s.title,
                s.creation_date,
                "; ".join(s.software),
                "; ".join(s.warnings),
            ]
            for i, val in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=i, value=val)
                cell.border = st["border"]
                cell.alignment = st["top_wrap"]
                if not s.author_found:
                    cell.fill = st["no_fill"]
            autor_cell = ws.cell(row=r, column=4)
            autor_cell.font = st["yes_font"] if s.author_found else st["no_font"]
            autor_cell.alignment = st["center"]

        for i, w in enumerate([38, 40, 14, 9, 45, 30, 20, 38, 45], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "B5"
        last_row = header_row + len(summaries)
        if summaries:
            ws.auto_filter.ref = f"A{header_row}:{last_col}{last_row}"
    else:
        # ---------- Hoja 1: Datos (según plantilla) ----------
        ws = wb.active
        ws.title = "Datos"
        _add_template_sheet(ws, summaries, template)

    # ---------- Hoja: Metadatos completos ----------
    _add_metadata_sheet(wb, summaries)

    # ---------- Hoja: Eliminados (si el usuario quitó archivos) ----------
    _add_deleted_sheet(wb, list(deleted or []))

    wb.save(path)
