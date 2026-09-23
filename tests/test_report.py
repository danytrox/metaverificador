import json

from openpyxl import load_workbook

from app.analyzer import Summary
from app.report import export_csv, export_html, export_json, export_xlsx
from app.template import load_template


def _summaries():
    s1 = Summary(
        path="/tmp/a.pdf", filename="a.pdf", filetype="PDF",
        author_found=True, author_values=["XMP-dc:Creator: Ana"],
        title="T1", creation_date="2024-01-01",
        software=["PDF:Creator: Word"], warnings=[],
        tags={"XMP-dc:Creator": "Ana"},
    )
    s2 = Summary(
        path="/tmp/b.pdf", filename="b.pdf", filetype="PDF",
        author_found=False, author_values=[],
        title="", creation_date="", software=[], warnings=["Sin autor"], tags={},
    )
    return [s1, s2]


def test_export_csv_default(tmp_path):
    p = tmp_path / "r.csv"
    export_csv(_summaries(), str(p))
    text = p.read_text(encoding="utf-8-sig")
    assert "Archivo" in text
    assert "SÍ" in text


def test_export_csv_con_plantilla(tmp_path):
    t = tmp_path / "t.html"
    t.write_text("<table><tr><th>Archivo</th><th>Autor</th></tr></table>", encoding="utf-8")
    spec = load_template(str(t))
    p = tmp_path / "r.csv"
    export_csv(_summaries(), str(p), template=spec)
    text = p.read_text(encoding="utf-8-sig")
    assert "Autor" in text
    assert "Ruta" not in text  # la plantilla no pide la columna Ruta


def test_export_html(tmp_path):
    p = tmp_path / "r.html"
    export_html(_summaries(), str(p))
    text = p.read_text(encoding="utf-8")
    assert "sin-autor" in text
    assert "<th>Archivo</th>" in text


def test_export_html_con_plantilla(tmp_path):
    t = tmp_path / "t.html"
    t.write_text("<table><tr><th>Autor</th></tr></table>", encoding="utf-8")
    spec = load_template(str(t))
    p = tmp_path / "r.html"
    export_html(_summaries(), str(p), template=spec)
    text = p.read_text(encoding="utf-8")
    assert "<th>Autor</th>" in text
    assert "<th>Tipo</th>" not in text


def test_export_json(tmp_path):
    p = tmp_path / "r.json"
    export_json(_summaries(), str(p))
    data = json.loads(p.read_text(encoding="utf-8"))
    assert len(data["files"]) == 2


def test_export_xlsx_default(tmp_path):
    p = tmp_path / "r.xlsx"
    export_xlsx(_summaries(), str(p))
    wb = load_workbook(str(p))
    assert "Resumen" in wb.sheetnames
    assert "Metadatos" in wb.sheetnames


def test_export_xlsx_con_plantilla(tmp_path):
    t = tmp_path / "t.html"
    t.write_text("<table><tr><th>Autor</th></tr></table>", encoding="utf-8")
    spec = load_template(str(t))
    p = tmp_path / "r.xlsx"
    export_xlsx(_summaries(), str(p), template=spec)
    wb = load_workbook(str(p))
    assert "Datos" in wb.sheetnames
